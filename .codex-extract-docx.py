from docx import Document
import sys


document = Document(sys.argv[1])

for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P{index}: {text}")

for table_index, table in enumerate(document.tables):
    print(f"TABLE {table_index}")
    for row in table.rows:
        values = [" ".join(cell.text.split()) for cell in row.cells]
        print(" | ".join(values))
